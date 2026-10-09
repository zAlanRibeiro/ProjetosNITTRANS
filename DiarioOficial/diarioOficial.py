# -*- coding: utf-8 -*-
"""
Junta numa planilha as portarias da NITTRANS publicadas no Diário Oficial de
Niterói no ano atual: Data / Portaria / Página / Conteúdo da Portaria.

O site não lista as edições: cada dia com edição tem um PDF em
do/{ano}/{MM_Mês}/{dia}.pdf e os demais dias dão 404. Por isso todos os dias
são tentados, de 1º de janeiro (ou do dia seguinte à última data que a
planilha já tem) até hoje.
"""
import io
import re
import time
import urllib.error
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

URL_EDICAO = "https://diariooficial.niteroi.rj.gov.br/do/{ano}/{mes:02d}_{nome_mes}/{dia:02d}.pdf"
NOMES_MESES = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun',
               'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']

# Texto de cada edição já lida. Edição publicada não muda, então só os dias
# novos (e os que ainda não tinham edição) são baixados de novo.
PASTA_CACHE = Path("cache")
PASTA_RESULTADOS = Path("resultados")

TENTATIVAS_DOWNLOAD = 3

COLUNAS = ['Data', 'Portaria', 'Página', 'Conteúdo da Portaria']
LARGURAS = [12, 12, 9, 120]
COR_NITTRANS = 'F97316'  # o mesmo laranja dos botões do Hub

# Cabeçalho: "PORTARIA NITTRANS nº 619/2026- O Presidente...". Também sai
# sem o ano ("nº 408-"), sem espaço ("nº340/2026") ou sozinho na linha.
# Citação no começo de linha ("Portaria NITTRANS nº 11/2022, publicada em")
# não é cabeçalho: depois do número só vem traço, ponto, fim de linha ou o
# texto em maiúscula.
RE_INICIO_NITTRANS = re.compile(
    r'^(?:PORTARIA|Portaria)\s+NITTRANS\s+[nN][º°ª.]{0,3}\s*(\d+)(?:/\s*(\d{4}))?'
    r'(?=\s*(?:[-–—.]|$|[A-Z"“]))[\s.]*[-–—]?\s*'
)
# Portaria de outro órgão logo depois: "Portaria SMA nº 264/2026 - ..."
RE_INICIO_OUTRA = re.compile(r'^(PORTARIA|Portaria|Port\.)\s.*\d/\d{4}\s*[-–—]')
# Última frase de toda portaria: "Esta Portaria entra(rá) em vigor ...".
RE_VIGOR = re.compile(r'Esta\s+Portaria\s+entra(?:rá)?\s+em\s+vigor[^.]*\.', re.I)
# Título de seção em caixa alta ("EXTRATO DE CONTRATO", nome de outro órgão):
# encerra a portaria publicada sem a frase de vigência.
RE_TITULO = re.compile(r'^(?!RESOLVE|CONSIDERANDO|ART)[^a-zà-ÿ“"]*[A-ZÀ-Ý]{2,}\s+[^a-zà-ÿ]*$')
# Rodapé de página: toda página termina em "Página N". Cai no meio das
# portarias longas e é por ele que se sabe em que página cada uma começa.
RE_RODAPE = re.compile(r'^Página\s+(\d+)$')
# Data no topo de cada página.
RE_DATA_TOPO = re.compile(r'^\d{2}/\d{2}/\d{4}$')
# Quebras de linha do PDF são só de diagramação; estas marcam parágrafos.
RE_PARAGRAFO = re.compile(r'\s+(?=Considerando\s|RESOLVE:|Art\.\s*\d)')


def extrair_portarias(texto, ano):
    """Devolve [(número, página, conteúdo)] das portarias NITTRANS de uma edição.

    'ano' completa o número das portarias publicadas sem o ano.
    """
    portarias = []
    atual = None
    pagina = 1
    for linha in texto.splitlines():
        linha = linha.strip()
        rodape = RE_RODAPE.match(linha)
        if rodape:
            pagina = int(rodape.group(1)) + 1
            continue
        if not linha or RE_DATA_TOPO.match(linha):
            continue
        inicio = RE_INICIO_NITTRANS.match(linha)
        if inicio:
            numero = f"{int(inicio.group(1))}/{inicio.group(2) or ano}"
            atual = [numero, pagina, [linha[inicio.end():]]]
            portarias.append(atual)
        elif RE_INICIO_OUTRA.match(linha) or RE_TITULO.match(linha):
            atual = None
        elif atual:
            atual[2].append(linha)

    resultado = []
    for numero, pagina, linhas in portarias:
        conteudo = ' '.join(linhas)
        # Depois da vigência vêm extratos, editais e outros órgãos.
        fim = RE_VIGOR.search(conteudo)
        if fim:
            conteudo = conteudo[:fim.end()]
        resultado.append((numero, pagina, RE_PARAGRAFO.sub('\n', conteudo).strip()))
    return resultado


def _baixar(url):
    """Bytes do PDF, ou None quando o dia não tem edição (404)."""
    for tentativa in range(1, TENTATIVAS_DOWNLOAD + 1):
        try:
            pedido = urllib.request.Request(url, headers={'User-Agent': 'HubNITTRANS'})
            with urllib.request.urlopen(pedido, timeout=60) as resposta:
                return resposta.read()
        except urllib.error.HTTPError as erro:
            if erro.code == 404:
                return None
            falha = erro
        except (urllib.error.URLError, TimeoutError, ConnectionError) as erro:
            falha = erro
        if tentativa < TENTATIVAS_DOWNLOAD:
            time.sleep(5 * tentativa)
    raise falha


def _url_edicao(dia):
    return URL_EDICAO.format(ano=dia.year, mes=dia.month,
                             nome_mes=NOMES_MESES[dia.month - 1], dia=dia.day)


def _texto_da_edicao(dia):
    """Texto da edição do dia, ou None se não houve edição."""
    cache = PASTA_CACHE / f"{dia:%Y-%m-%d}.txt"
    if cache.exists():
        return cache.read_text(encoding='utf-8')

    try:
        pdf_bytes = _baixar(_url_edicao(dia))
    except Exception as erro:
        # Parar é melhor que entregar a planilha sem um dia sem ninguém notar.
        # O que já foi lido fica no cache, então rodar de novo é rápido.
        raise RuntimeError(
            f"Não foi possível baixar o Diário Oficial de {dia:%d/%m/%Y}"
            f" ({erro}). Rode de novo: os dias já lidos ficam guardados."
        ) from erro
    if pdf_bytes is None:
        return None

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        texto = '\n'.join((pagina.extract_text() or '') for pagina in pdf.pages)
    cache.write_text(texto, encoding='utf-8')
    return texto


def _nova_planilha():
    livro = Workbook()
    aba = livro.active
    aba.title = 'Portarias'
    aba.append(COLUNAS)
    for celula, largura in zip(aba[1], LARGURAS):
        celula.fill = PatternFill('solid', fgColor=COR_NITTRANS)
        celula.font = Font(bold=True, color='FFFFFF')
        aba.column_dimensions[celula.column_letter].width = largura
    aba.freeze_panes = 'A2'
    return livro


def _posicao_das_colunas(aba, caminho):
    """Posição de cada coluna pelo nome, para não escrever na coluna errada
    se alguém mudou a ordem ou incluiu colunas na planilha."""
    cabecalho = {celula.value: celula.column for celula in aba[1]}
    faltando = [nome for nome in COLUNAS if nome not in cabecalho]
    if faltando:
        raise RuntimeError(
            f"A planilha {caminho.name} não tem a(s) coluna(s)"
            f" {', '.join(faltando)}. Volte o nome original da coluna ou"
            " apague a planilha para ela ser gerada de novo."
        )
    return cabecalho


def rodar_diario_oficial():
    for pasta in [PASTA_CACHE, PASTA_RESULTADOS]:
        pasta.mkdir(exist_ok=True)

    hoje = date.today()
    caminho = PASTA_RESULTADOS / f"Portarias_NITTRANS_{hoje.year}.xlsx"

    # Planilha já preenchida: busca só depois da última data dela, e o que
    # já está lá (inclusive anotações feitas à mão) fica como está.
    inicio = date(hoje.year, 1, 1)
    if caminho.exists():
        livro = load_workbook(caminho)
        aba = livro['Portarias']
        colunas = _posicao_das_colunas(aba, caminho)
        datas = [celula.value for (celula,) in aba.iter_rows(
            min_row=2, min_col=colunas['Data'], max_col=colunas['Data'])
            if isinstance(celula.value, datetime)]
        if datas:
            inicio = max(datas).date() + timedelta(days=1)
    else:
        livro = _nova_planilha()
        aba = livro['Portarias']
        colunas = _posicao_das_colunas(aba, caminho)

    novas = 0
    dia = inicio
    while dia <= hoje:
        texto = _texto_da_edicao(dia)
        for numero, pagina, conteudo in extrair_portarias(texto or '', dia.year):
            linha = aba.max_row + 1
            aba.cell(linha, colunas['Data'], dia).number_format = 'DD/MM/YYYY'
            aba.cell(linha, colunas['Portaria'], numero)
            celula = aba.cell(linha, colunas['Página'], pagina)
            celula.hyperlink = f"{_url_edicao(dia)}#page={pagina}"
            celula.style = 'Hyperlink'
            aba.cell(linha, colunas['Conteúdo da Portaria'], conteudo).alignment = (
                Alignment(wrap_text=True, vertical='top'))
            novas += 1
        dia += timedelta(days=1)

    aba.auto_filter.ref = aba.dimensions
    try:
        livro.save(caminho)
    except PermissionError:
        raise RuntimeError(
            f"Feche a planilha {caminho.name} no Excel e clique de novo."
        ) from None

    periodo = f"{inicio:%d/%m/%Y} a {hoje:%d/%m/%Y}" if inicio <= hoje else "nada novo"
    print(f"\nConcluído! {novas} portaria(s) nova(s) ({periodo}) -> {caminho}")


if __name__ == "__main__":
    rodar_diario_oficial()
