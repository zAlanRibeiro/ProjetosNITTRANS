# -*- coding: utf-8 -*-
"""
Junta os CSVs do relatório de demandas e atendimentos do Colab numa única
planilha Excel, com uma aba por recorte.

O Colab exporta o relatório picado em seis CSVs
(colab-demands-and-attendance-report-data-0 a -5), cada um com um recorte
diferente. Esta ferramenta identifica cada arquivo pelo número no nome —
não pela ordem em que aparecem na pasta — trata os dados e traduz colunas
e valores para português.

O que cada arquivo traz (conferido contra os dados reais):
  0  data/hora e plataforma de cada demanda (a base, uma linha por demanda)
  1  total por status e categoria
  2  total por faixa de tempo de atendimento e categoria
  3  total e taxa de resolução por bairro, plataforma Colab
  4  total e taxa de resolução por bairro, plataforma cdo
  5  total por status e bairro

Os arquivos 3 e 4 têm colunas idênticas e só se distinguem pela
plataforma: a soma de cada um bate exatamente com a contagem por
plataforma do arquivo 0 (5.809 Colab e 2.566 cdo na amostra analisada).

Tratamento aplicado (tudo o que for corrigido é listado no terminal,
nunca corrigido em silêncio):
  - espaços sobrando nas pontas e espaços repetidos no meio dos textos
  - números e datas conferidos um a um; linha com valor ilegível é
    descartada e reportada, em vez de virar zero ou data vazia
"""

import re
import shutil
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PASTA_ENTRADA = Path("entrada")
PASTA_BACKUP = Path("backup")
PASTA_RESULTADOS = Path("resultados")

# Vai no fim do nome da planilha. Como os seis CSVs viram um arquivo só,
# o nome não sai de nenhum deles em particular.
NOME_SAIDA = "Colab_Demandas_Extraido"

# Os horários do Colab vêm em UTC (terminados em Z). Sem converter, uma
# demanda registrada às 21h de Niterói cairia no dia seguinte na planilha.
FUSO_LOCAL = "America/Sao_Paulo"

# ==========================================
# 1. Tradução de colunas e valores
# ==========================================
STATUS = {'open': 'Aberta', 'solved': 'Resolvida'}

# A ordem importa: a faixa de tempo é uma escala, não uma lista alfabética.
FAIXAS_TEMPO = [
    ('7 days or less', 'Até 7 dias'),
    ('7-15 days', '7 a 15 dias'),
    ('15 days - 3 months', '15 dias a 3 meses'),
    ('3-6 months', '3 a 6 meses'),
    ('6-12 months', '6 a 12 meses'),
    ('More than 12 months', 'Mais de 12 meses'),
]
TEMPO = dict(FAIXAS_TEMPO)
ORDEM_TEMPO = [pt for _, pt in FAIXAS_TEMPO]

COLUNAS = {
    'createdAt': 'Data/Hora',
    'platform': 'Plataforma',
    'total': 'Total',
    'status': 'Status',
    'category': 'Categoria',
    'time': 'Tempo de Atendimento',
    'totals_per_category': 'Total da Categoria',
    'neighborhood': 'Bairro',
    'rate': 'Taxa de Resolução',
}

# ==========================================
# 2. Definição das abas
# ==========================================
# numero do arquivo -> (nome da aba, colunas esperadas no CSV)
ABAS = {
    0: ('Demandas', ['createdAt', 'platform']),
    1: ('Status por Categoria', ['total', 'status', 'category']),
    2: ('Tempo por Categoria',
        ['total', 'time', 'category', 'totals_per_category']),
    3: ('Bairros (Colab)', ['neighborhood', 'total', 'rate']),
    4: ('Bairros (cdo)', ['neighborhood', 'total', 'rate']),
    5: ('Status por Bairro', ['total', 'status', 'neighborhood']),
}

COLUNAS_INTEIRAS_CSV = {'total', 'totals_per_category'}
COLUNAS_DECIMAIS_CSV = {'rate'}
COLUNAS_DATA_CSV = {'createdAt'}

COLUNAS_PERCENTUAL = {'Taxa de Resolução'}
COLUNAS_INTEIRAS = {'Total', 'Total da Categoria', 'Resolvidas'}

RE_NUMERO_ARQUIVO = re.compile(r'data-(\d+)', re.IGNORECASE)


def _numero_do_arquivo(caminho):
    """
    Descobre qual dos seis recortes é o arquivo, pelo número no nome.
    O Colab acrescenta sufixos de download (' (1)', ' (2)'), por isso a
    identificação não pode depender do nome inteiro nem da ordem da pasta.
    """
    encontrado = RE_NUMERO_ARQUIVO.search(caminho.name)
    return int(encontrado.group(1)) if encontrado else None


# ==========================================
# 3. Tratamento dos dados
# ==========================================
def _limpar_textos(df, numero, avisos):
    """
    Tira espaços das pontas e junta espaços repetidos no meio. O Colab
    exporta categorias como 'Evento irregular ' (com espaço no fim), que
    no Excel viram uma categoria diferente de 'Evento irregular'.
    """
    for coluna in df.columns:
        if (coluna in COLUNAS_INTEIRAS_CSV or coluna in COLUNAS_DECIMAIS_CSV
                or coluna in COLUNAS_DATA_CSV):
            continue
        # Em pandas 3 as colunas de texto não são 'object'; converter para
        # 'string' garante que o tratamento roda em todas elas.
        original = df[coluna].astype('string')
        limpo = original.str.replace(r'\s+', ' ', regex=True).str.strip()
        mudou = original.notna() & (limpo != original)
        if mudou.any():
            exemplos = sorted(set(original[mudou]))[:5]
            avisos.append(
                (numero, f"{int(mudou.sum())} valor(es) de '{coluna}' com"
                         f" espaços sobrando, corrigidos: "
                         + ', '.join(repr(e) for e in exemplos))
            )
        df[coluna] = limpo
    return df


def _converter_numeros(df, numero, avisos):
    """
    Converte os números conferindo um a um. Linha com valor ilegível é
    descartada e reportada — virar zero calado falsearia os totais.
    """
    ruins = pd.Series(False, index=df.index)

    for coluna in df.columns:
        if coluna in COLUNAS_INTEIRAS_CSV:
            valores = pd.to_numeric(df[coluna], errors='coerce')
            invalido = df[coluna].notna() & valores.isna()
            invalido |= valores.notna() & ((valores < 0) | (valores % 1 != 0))
        elif coluna in COLUNAS_DECIMAIS_CSV:
            valores = pd.to_numeric(df[coluna], errors='coerce')
            invalido = df[coluna].notna() & valores.isna()
            invalido |= valores.notna() & ((valores < 0) | (valores > 1))
        else:
            continue

        vazio = df[coluna].isna()
        if vazio.any():
            avisos.append((numero, f"{int(vazio.sum())} linha(s) sem"
                                   f" '{coluna}' — descartadas"))
        if invalido.any():
            exemplos = sorted(set(df.loc[invalido, coluna].dropna()))[:5]
            avisos.append(
                (numero, f"{int(invalido.sum())} linha(s) com '{coluna}'"
                         f" fora do esperado — descartadas: "
                         + ', '.join(repr(e) for e in exemplos))
            )
        ruins |= invalido | vazio
        df[coluna] = valores

    if ruins.any():
        df = df[~ruins].reset_index(drop=True)
    for coluna in df.columns:
        if coluna in COLUNAS_INTEIRAS_CSV:
            df[coluna] = df[coluna].astype(int)
    return df


def _converter_datas(df, numero, avisos):
    """Descarta e reporta data ilegível, em vez de deixar célula vazia."""
    if 'createdAt' not in df.columns:
        return df
    momento = pd.to_datetime(df['createdAt'], format='ISO8601',
                             utc=True, errors='coerce')
    invalido = momento.isna()
    if invalido.any():
        exemplos = sorted(set(df.loc[invalido, 'createdAt'].dropna()))[:5]
        avisos.append(
            (numero, f"{int(invalido.sum())} linha(s) com data/hora ilegível"
                     f" — descartadas: "
                     + ', '.join(repr(e) for e in exemplos))
        )
    df = df[~invalido].copy()
    df['createdAt'] = momento[~invalido].dt.tz_convert(
        FUSO_LOCAL).dt.tz_localize(None)
    return df.reset_index(drop=True)


def tratar(df, numero, avisos):
    df = _limpar_textos(df, numero, avisos)
    df = _converter_datas(df, numero, avisos)
    df = _converter_numeros(df, numero, avisos)
    return df


# ==========================================
# 4. Montagem de cada aba
# ==========================================
def _preparar_demandas(df):
    """Arquivo 0: uma linha por demanda, com data e hora separadas."""
    return pd.DataFrame({
        'Data': df['createdAt'].dt.date,
        'Hora': df['createdAt'].dt.strftime('%H:%M:%S'),
        'Plataforma': df['platform'],
    }).sort_values(['Data', 'Hora']).reset_index(drop=True)


def _preparar_bairros(df):
    """
    Arquivos 3 e 4: acrescenta a quantidade de demandas resolvidas, que é
    o que a taxa representa. total * rate dá inteiro exato em todas as
    linhas, então a coluna é o dado original, não uma estimativa.
    """
    saida = df.rename(columns=COLUNAS)
    saida['Resolvidas'] = (
        saida['Total'] * saida['Taxa de Resolução']).round().astype(int)
    saida = saida[['Bairro', 'Total', 'Resolvidas', 'Taxa de Resolução']]
    return saida.sort_values('Total', ascending=False).reset_index(drop=True)


def _preparar_tempo(df):
    """Arquivo 2: faixa de tempo em ordem cronológica, não alfabética."""
    saida = df.rename(columns=COLUNAS)
    saida['Tempo de Atendimento'] = pd.Categorical(
        saida['Tempo de Atendimento'].map(TEMPO).fillna(
            saida['Tempo de Atendimento']),
        categories=ORDEM_TEMPO, ordered=True,
    )
    saida = saida[['Categoria', 'Tempo de Atendimento', 'Total',
                   'Total da Categoria']]
    return saida.sort_values(
        ['Total da Categoria', 'Categoria', 'Tempo de Atendimento'],
        ascending=[False, True, True],
    ).reset_index(drop=True)


def _preparar_por_status(df, coluna_grupo):
    """Arquivos 1 e 5: total por status e por categoria/bairro."""
    saida = df.rename(columns=COLUNAS)
    saida['Status'] = saida['Status'].map(STATUS).fillna(saida['Status'])
    saida = saida[[coluna_grupo, 'Status', 'Total']]
    return saida.sort_values([coluna_grupo, 'Status']).reset_index(drop=True)


def preparar(numero, df):
    if numero == 0:
        return _preparar_demandas(df)
    if numero == 1:
        return _preparar_por_status(df, 'Categoria')
    if numero == 2:
        return _preparar_tempo(df)
    if numero in (3, 4):
        return _preparar_bairros(df)
    if numero == 5:
        return _preparar_por_status(df, 'Bairro')
    return df.rename(columns=COLUNAS)


# ==========================================
# 5. Geração da planilha
# ==========================================
COR_CABECALHO = '1F4E78'


def _caminho_livre(caminho):
    """
    Nunca sobrescreve uma planilha já existente: se o nome estiver ocupado,
    acrescenta ' (2)', ' (3)'... como o Windows faz.
    """
    if not caminho.exists():
        return caminho
    contador = 2
    while True:
        candidato = caminho.with_name(
            f'{caminho.stem} ({contador}){caminho.suffix}'
        )
        if not candidato.exists():
            return candidato
        contador += 1


def _formatar_aba(ws, colunas):
    fonte_cabecalho = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    fill_cabecalho = PatternFill(
        start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type='solid'
    )
    borda_fina = Side(border_style='thin', color='D3D3D3')
    borda = Border(left=borda_fina, right=borda_fina,
                   top=borda_fina, bottom=borda_fina)

    for indice, nome in enumerate(colunas, start=1):
        letra = get_column_letter(indice)
        celula = ws.cell(row=1, column=indice)
        celula.font = fonte_cabecalho
        celula.fill = fill_cabecalho
        celula.alignment = Alignment(
            horizontal='center', vertical='center', wrap_text=True
        )
        celula.border = borda

        if nome == 'Data':
            formato, largura = 'DD/MM/YYYY', 14
        elif nome in COLUNAS_PERCENTUAL:
            formato, largura = '0.0%', 16
        elif nome in COLUNAS_INTEIRAS:
            formato, largura = '#,##0', 14
        else:
            formato, largura = None, max(len(nome) + 4, 14)

        ws.column_dimensions[letra].width = largura
        if formato:
            for linha in range(2, ws.max_row + 1):
                ws.cell(row=linha, column=indice).number_format = formato

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def _gerar_excel(tabelas, caminho_saida):
    """Uma aba por recorte, na ordem dos arquivos. Nada além disso."""
    with pd.ExcelWriter(caminho_saida, engine='openpyxl') as writer:
        for numero in sorted(tabelas):
            aba = ABAS[numero][0][:31]
            tabelas[numero].to_excel(writer, index=False, sheet_name=aba)
            _formatar_aba(writer.sheets[aba], list(tabelas[numero].columns))


# ==========================================
# 6. Função principal
# ==========================================
def rodar_colab_demandas():
    for pasta in [PASTA_ENTRADA, PASTA_BACKUP, PASTA_RESULTADOS]:
        pasta.mkdir(exist_ok=True)

    arquivos_csv = sorted(PASTA_ENTRADA.glob('*.csv'))
    if not arquivos_csv:
        print(f"Nenhum CSV encontrado em '{PASTA_ENTRADA}'.")
        return

    tabelas = {}
    ignorados = []
    avisos = []

    for arquivo in arquivos_csv:
        numero = _numero_do_arquivo(arquivo)
        if numero is None or numero not in ABAS:
            ignorados.append((arquivo.name, 'não é um dos seis recortes'
                                            ' esperados (data-0 a data-5)'))
            continue

        # Tudo entra como texto: a conversão de número e data é feita
        # depois, com conferência, em vez de confiar no palpite do pandas.
        df = pd.read_csv(arquivo, encoding='utf-8-sig', dtype=str)
        faltando = [c for c in ABAS[numero][1] if c not in df.columns]
        if faltando:
            ignorados.append((arquivo.name,
                              'faltam colunas: ' + ', '.join(faltando)))
            continue

        if numero in tabelas:
            ignorados.append((arquivo.name,
                              f'já havia um arquivo data-{numero} nesta'
                              ' execução'))
            continue

        linhas_lidas = len(df)
        df = tratar(df, numero, avisos)
        tabelas[numero] = preparar(numero, df)

        print(f'  {arquivo.name}')
        perdidas = linhas_lidas - len(tabelas[numero])
        sufixo = f' ({perdidas} descartada(s))' if perdidas else ''
        print(f'    -> aba "{ABAS[numero][0]}":'
              f' {len(tabelas[numero])} linhas{sufixo}')

    if avisos:
        print('\n' + '=' * 60)
        print('DADOS TRATADOS (conferir)')
        print('=' * 60)
        for numero, mensagem in avisos:
            print(f'  [{ABAS[numero][0]}] {mensagem}')
        print('=' * 60)

    if ignorados:
        print('\n' + '=' * 60)
        print('ARQUIVOS IGNORADOS')
        print('=' * 60)
        for nome, motivo in ignorados:
            print(f'  {nome}: {motivo}')
        print('=' * 60)

    if not tabelas:
        print('\nNenhum CSV do relatório do Colab foi reconhecido.')
        return

    ausentes = [f'data-{n} ({ABAS[n][0]})' for n in ABAS if n not in tabelas]
    if ausentes:
        print('\n' + '=' * 60)
        print('RECORTES QUE NÃO VIERAM (a planilha sai sem essas abas)')
        print('=' * 60)
        for nome in ausentes:
            print(f'  {nome}')
        print('=' * 60)

    caminho_saida = _caminho_livre(PASTA_RESULTADOS / f'{NOME_SAIDA}.xlsx')
    _gerar_excel(tabelas, caminho_saida)

    for arquivo in arquivos_csv:
        shutil.move(str(arquivo), str(PASTA_BACKUP / arquivo.name))

    print('\n' + '=' * 60)
    print(f'Concluído! Planilha: {caminho_saida}')
    print(f'  {len(tabelas)} de {len(ABAS)} abas.')
    print('=' * 60)


if __name__ == '__main__':
    rodar_colab_demandas()
