# -*- coding: utf-8 -*-
import os
import sys
import threading
import tkinter as tk
from tkinter import filedialog, messagebox
from PIL import Image, ImageDraw, ImageTk
import customtkinter as ctk
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def bloquear_e_formatar_planilha(
    caminho_entrada, caminho_saida, senha, modo_bloqueio="total"
):
  """Abre uma planilha, aplica formatação profissional e protege com senha.

  modo_bloqueio:
    "total"   -> bloqueia a aba inteira (comportamento padrão/atual).
    "parcial" -> reservado para uma futura opção de bloqueio parcial.
  """
  wb = openpyxl.load_workbook(caminho_entrada)
  ws = wb.active

  # Garante que as linhas de grade estejam visíveis
  ws.views.sheetView[0].showGridLines = True

  # Ativa oficialmente a proteção por senha na aba (Bloqueio Total)
  ws.protection.sheet = True
  ws.protection.password = senha

  # Estilos Visuais Profissionais
  cor_cabecalho = "1F4E78"
  fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
  fill_cabecalho = PatternFill(
      start_color=cor_cabecalho, end_color=cor_cabecalho, fill_type="solid"
  )

  borda_fina = Side(border_style="thin", color="D3D3D3")
  borda_celula = Border(
      left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina
  )

  for row in range(1, ws.max_row + 1):
    for col in range(1, ws.max_column + 1):
      celula = ws.cell(row=row, column=col)
      celula.border = borda_celula

      # Primeira linha é o cabeçalho: fundo azul, texto branco em negrito
      if row == 1:
        celula.font = fonte_cabecalho
        celula.fill = fill_cabecalho
        celula.alignment = Alignment(horizontal="center", vertical="center")
        continue

      celula.font = Font(name="Calibri", size=11)

      if isinstance(celula.value, (int, float)):
        celula.alignment = Alignment(horizontal="right", vertical="center")
      else:
        celula.alignment = Alignment(horizontal="left", vertical="center")

  # Ajustando automaticamente a largura das colunas
  for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
      if cell.value:
        max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

  wb.save(caminho_saida)
  return caminho_saida


# ── Paleta de Cores IDÊNTICA ao Hub e Criptografia ──────────────────────────
COR_VIDRO_BLENDED = "#0A2548"
_AZUL_SUB = "#A3C2F0"
_LARANJA = "#F97316"
_LAR_HOV = "#EA580C"
_BRANCO = "#FFFFFF"
_BORDA = "#E2E8F0"
_TEXTO = "#1A202C"
_TEXTO_S = "#A3C2F0"
_SUCESSO = "#4ADE80"
_ERRO = "#F87171"
_FONTE = "Segoe UI"

TIPOS_ABRIR = [
    ("Planilhas Excel", "*.xlsx *.xls"),
    ("Todos os arquivos", "*.*"),
]


class _AppMixin:
  """Lógica e UI compartilhadas para a janela de Bloquear Planilha."""

  def _init_common(self):
    self.title("Bloquear Planilha — NITTRANS")
    self.geometry("480x530")
    self.resizable(False, False)
    self.configure(fg_color=COR_VIDRO_BLENDED)

    self.after(50, self._aplicar_icone)
    self.after(400, self._aplicar_icone)

    self._caminho = ""
    self._after_progresso = None
    self._build_ui()
    self._centralizar()

  def _obter_caminho_recurso(self, nome_arquivo):
    caminhos = [
        nome_arquivo,
        os.path.join("..", nome_arquivo),
        os.path.join(os.path.dirname(__file__), "..", nome_arquivo),
        os.path.join(getattr(sys, "_MEIPASS", ""), nome_arquivo),
        os.path.join(
            os.path.dirname(sys.executable), "_internal", nome_arquivo
        ),
    ]
    for p in caminhos:
      if os.path.exists(p):
        return p
    return nome_arquivo

  def _aplicar_icone(self):
    try:
      ico = self._obter_caminho_recurso("logo.ico")
      if os.path.exists(ico):
        self.iconbitmap(ico)
    except Exception:
      pass

  def _centralizar(self):
    self.update_idletasks()
    w, h = self.winfo_width(), self.winfo_height()
    sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
    self.geometry(f"{w}x{h}+{(sw - w) // 2}+{(sh - h) // 2}")

  def _build_ui(self):
    self.canvas = tk.Canvas(
        self,
        width=480,
        height=530,
        bg=COR_VIDRO_BLENDED,
        highlightthickness=0,
    )
    self.canvas.place(x=0, y=0)

    caminho_fundo = self._obter_caminho_recurso("FundoAplicativoNITTRANS.png")
    if os.path.exists(caminho_fundo):
      try:
        img_original = Image.open(caminho_fundo).convert("RGBA")
        img_original = img_original.resize((480, 530), Image.LANCZOS)
        camada_overlay = Image.new("RGBA", img_original.size, (0, 0, 0, 0))
        draw = ImageDraw.Draw(camada_overlay)
        draw.rounded_rectangle(
            (15, 15, 465, 515), radius=15, fill=(8, 25, 55, 110)
        )
        img_mesclada = Image.alpha_composite(img_original, camada_overlay)
        self.tk_fundo = ImageTk.PhotoImage(img_mesclada)
        self.canvas.create_image(0, 0, anchor="nw", image=self.tk_fundo)
      except Exception as e:
        print(f"Erro ao carregar fundo: {e}")

    self.canvas.create_text(
        40,
        45,
        text="Bloquear & Formatar Planilha",
        font=(_FONTE, 18, "bold"),
        fill=_BRANCO,
        anchor="w",
    )
    self.canvas.create_text(
        40,
        70,
        text="Proteja planilhas Excel com senha e layout profissional",
        font=(_FONTE, 11),
        fill=_AZUL_SUB,
        anchor="w",
    )

    self.canvas.create_text(
        40, 115, text="Arquivo", font=(_FONTE, 12, "bold"), fill=_BRANCO, anchor="w"
    )
    self.canvas.create_text(
        40,
        195,
        text="Senha de proteção",
        font=(_FONTE, 12, "bold"),
        fill=_BRANCO,
        anchor="w",
    )
    self.canvas.create_text(
        40,
        275,
        text="Confirmar senha",
        font=(_FONTE, 12, "bold"),
        fill=_BRANCO,
        anchor="w",
    )

    self.id_status = self.canvas.create_text(
        240,
        490,
        text="Selecione uma planilha para começar.",
        font=(_FONTE, 11),
        fill=_TEXTO_S,
        anchor="center",
    )

    # Widgets Flutuantes
    self.entry_arquivo = ctk.CTkEntry(
        self,
        placeholder_text="Nenhum arquivo selecionado",
        state="disabled",
        height=40,
        width=280,
        corner_radius=8,
        fg_color=_BRANCO,
        border_color=_BORDA,
        text_color=_TEXTO,
        bg_color=COR_VIDRO_BLENDED,
    )
    self.entry_arquivo.place(x=40, y=130)

    ctk.CTkButton(
        self,
        text="Selecionar",
        width=110,
        height=40,
        fg_color=_LARANJA,
        hover_color=_LAR_HOV,
        text_color="white",
        font=(_FONTE, 12, "bold"),
        corner_radius=8,
        bg_color=COR_VIDRO_BLENDED,
        command=self._selecionar,
    ).place(x=330, y=130)

    self.var_senha = tk.StringVar()
    self.entry_senha = ctk.CTkEntry(
        self,
        textvariable=self.var_senha,
        width=310,
        show="●",
        placeholder_text="Digite a senha",
        height=40,
        corner_radius=8,
        fg_color=_BRANCO,
        border_color=_BORDA,
        text_color=_TEXTO,
        bg_color=COR_VIDRO_BLENDED,
    )
    self.entry_senha.place(x=40, y=210)
    self.entry_senha.bind("<Return>", lambda e: self._processar())

    self.var_mostrar = tk.BooleanVar()
    ctk.CTkCheckBox(
        self,
        text="Mostrar",
        variable=self.var_mostrar,
        command=self._toggle_senha,
        width=80,
        fg_color=_LARANJA,
        hover_color=_LAR_HOV,
        bg_color=COR_VIDRO_BLENDED,
        text_color=_BRANCO,
        checkmark_color="white",
    ).place(x=365, y=218)

    self.var_confirmar = tk.StringVar()
    self.entry_confirmar = ctk.CTkEntry(
        self,
        textvariable=self.var_confirmar,
        width=310,
        show="●",
        placeholder_text="Confirme a senha",
        height=40,
        corner_radius=8,
        fg_color=_BRANCO,
        border_color=_BORDA,
        text_color=_TEXTO,
        bg_color=COR_VIDRO_BLENDED,
    )
    self.entry_confirmar.place(x=40, y=290)
    self.entry_confirmar.bind("<Return>", lambda e: self._processar())

    ctk.CTkButton(
        self,
        text="🔒   Bloquear Planilha",
        width=400,
        fg_color=_LARANJA,
        hover_color=_LAR_HOV,
        text_color="white",
        height=48,
        corner_radius=8,
        font=(_FONTE, 14, "bold"),
        bg_color=COR_VIDRO_BLENDED,
        command=self._processar,
    ).place(x=40, y=365)

    self.progress = ctk.CTkProgressBar(
        self,
        mode="determinate",
        width=400,
        progress_color=_LARANJA,
        bg_color=COR_VIDRO_BLENDED,
        fg_color=_BRANCO,
        height=6,
        corner_radius=3,
    )
    self.progress.place(x=40, y=450)
    self.progress.set(0)

  def _iniciar_progresso(self):
    self.progress.set(0)
    self._progresso_val = 0.0
    self._avancar_progresso()

  def _avancar_progresso(self):
    if self._progresso_val < 0.85:
      self._progresso_val = min(self._progresso_val + 0.018, 0.85)
      self.progress.set(self._progresso_val)
      self._after_progresso = self.after(30, self._avancar_progresso)

  def _finalizar_progresso(self, mensagem: str):
    if self._after_progresso:
      self.after_cancel(self._after_progresso)
      self._after_progresso = None
    self.progress.set(1.0)
    self.canvas.itemconfig(self.id_status, text=mensagem, fill=_SUCESSO)
    self.after(3000, lambda: self.progress.set(0))

  def _erro_progresso(self, mensagem: str):
    if self._after_progresso:
      self.after_cancel(self._after_progresso)
      self._after_progresso = None
    self.progress.set(0)
    self.canvas.itemconfig(self.id_status, text=mensagem, fill=_ERRO)

  def _toggle_senha(self):
    caractere = "" if self.var_mostrar.get() else "●"
    self.entry_senha.configure(show=caractere)
    self.entry_confirmar.configure(show=caractere)

  def _set_arquivo(self, caminho: str):
    self._caminho = caminho
    self.entry_arquivo.configure(state="normal")
    self.entry_arquivo.delete(0, "end")
    self.entry_arquivo.insert(0, os.path.basename(caminho))
    self.entry_arquivo.configure(state="disabled")

  def _selecionar(self):
    caminho = filedialog.askopenfilename(
        title="Selecione a planilha Excel", filetypes=TIPOS_ABRIR
    )
    if caminho:
      self._set_arquivo(caminho)
      self.canvas.itemconfig(
          self.id_status,
          text=f"Arquivo: {os.path.basename(caminho)}",
          fill=_TEXTO_S,
      )

  def _validar(self):
    if not self._caminho:
      messagebox.showwarning("Aviso", "Selecione uma planilha.")
      return None
    if not os.path.isfile(self._caminho):
      messagebox.showerror("Erro", "Arquivo não encontrado.")
      return None

    senha = self.var_senha.get()
    confirmar = self.var_confirmar.get()

    if not senha:
      messagebox.showwarning("Aviso", "Digite uma senha de proteção.")
      return None

    if senha != confirmar:
      messagebox.showwarning(
          "Aviso", "As senhas digitadas não coincidem. Verifique novamente."
      )
      return None

    return self._caminho, senha

  def _resetar_campos(self):
    self._caminho = ""
    self.entry_arquivo.configure(state="normal")
    self.entry_arquivo.delete(0, "end")
    self.entry_arquivo.configure(state="disabled")
    self.var_senha.set("")
    self.var_confirmar.set("")

  def _processar(self):
    r = self._validar()
    if not r:
      return
    caminho, senha = r

    nome_sugerido = "protegido_" + os.path.basename(caminho)
    destino = filedialog.asksaveasfilename(
        title="Salvar planilha protegida como",
        initialfile=nome_sugerido,
        defaultextension=".xlsx",
        filetypes=[("Planilha Excel", "*.xlsx"), ("Todos os arquivos", "*.*")],
    )
    if not destino:
      return

    self._iniciar_progresso()
    self.canvas.itemconfig(
        self.id_status, text="Aplicando bloqueio…", fill=_TEXTO_S
    )

    def tarefa():
          try:
            bloquear_e_formatar_planilha(caminho, destino, senha)
            nome = os.path.basename(destino)
            self.after(
                0, lambda: self._finalizar_progresso(f"✔ Finalizado: {nome}")
            )
            self.after(0, self._resetar_campos)
            self.after(
                50,
                lambda: messagebox.showinfo(
                    "Sucesso",
                    f"Planilha bloqueada com sucesso!\n\nSalva em:\n  {destino}",
                ),
            )
          except Exception as err:
            err_msg = str(err)
            self.after(
                0, lambda: self._erro_progresso("✖ Erro ao proteger planilha.")
            )
            self.after(
                50, lambda: messagebox.showerror("Erro no Processamento", err_msg)
            )


class AppToplevel(_AppMixin, ctk.CTkToplevel):
  """Janela filha — integrada na seção de segurança do Hub."""

  def __init__(self, master):
    ctk.CTkToplevel.__init__(self, master)
    self._init_common()


def rodar_bloqueio():
  """Método de compatibilidade caso o Hub tente chamá-lo diretamente."""
  pass